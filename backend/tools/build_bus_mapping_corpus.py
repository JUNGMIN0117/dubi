"""공식 서울 버스 XLSX와 로컬 카카오 응답으로 오프라인 검증 입력을 만든다.

운영 적재기를 대체하지 않는 테스트용 변환기다. OA-1095의 공식 ROUTE_ID/NODE_ID,
순번, WGS84 좌표를 유지하며 이름으로 노선 후보를 미리 제거하지 않는다. 모든 노선을
남겨 동명 노선과 순환 구간의 모호성을 실제 Resolver가 판단하게 한다.

입력은 읽기 전용이며 결과는 지정한 새 JSON 파일 하나다. API/DB/Redis에 접근하지 않는다.
지하철 운행 순서는 이 파일로 증명할 수 없으므로 subwayLines를 비워 명시하고, 정답 labels도
자동 생성하지 않는다. 파일 수정시각을 수집시각으로 사용하지 않으며 미상은 null로 유지한다.
필수 열·ID·순번·좌표 오류와 중복 노선 순번은 원본 Excel 행 번호와 함께 실패시킨다.
"""

import argparse
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


def identifier(value):
    """Excel 숫자 ID의 소수점 표기를 제거하되 누락이나 임의 반올림은 허용하지 않는다."""
    if isinstance(value, bool) or value is None:
        raise ValueError("ID 누락 또는 잘못된 타입")
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            raise ValueError("정수 ID 필요")
        value = int(value)
    result = str(value).strip()
    if not result:
        raise ValueError("빈 ID")
    return result


def read_bus_routes(path, route_ids=None):
    """공식 마스터 전체를 routeId별로 묶고 순서를 고정해 Replay 가능한 읽기 모델로 반환한다.

    반환 ID는 외부 카카오 ID가 아닌 서울 버스 공식 ID다. ARS_ID는 마스터 열 존재만
    확인하며 내부 canonical NODE_ID를 대체하지 않는다. 실패한 행을 조용히 제외하지 않는다.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows)
        required = ("ROUTE_ID", "노선명", "순번", "NODE_ID", "ARS_ID", "정류소명", "X좌표", "Y좌표")
        if any(header.count(name) != 1 for name in required):
            raise ValueError("공식 버스 마스터의 필수 열 누락 또는 중복")
        columns = {name: header.index(name) for name in required}
        routes, seen = {}, set()
        # 1. 행별 품질 검사. 오류 위치는 Excel 행 번호로 남겨 원본을 다시 확인할 수 있다.
        for row_number, row in enumerate(rows, 2):
            if all(value is None for value in row):
                continue
            try:
                route_id = identifier(row[columns["ROUTE_ID"]])
                # 명시한 공식 노선의 제한 실험만 허용한다. 이름 검색으로 임의 선택하지 않으며,
                # 사용한 범위는 masterSource에 기록되어 전체 마스터 coverage와 구별된다.
                if route_ids and route_id not in route_ids:
                    continue
                name = identifier(row[columns["노선명"]])
                sequence = int(identifier(row[columns["순번"]]))
                longitude, latitude = (float(row[columns[key]]) for key in ("X좌표", "Y좌표"))
                if sequence < 1 or not (-180 <= longitude <= 180 and -90 <= latitude <= 90):
                    raise ValueError("순번 또는 좌표 범위 오류")
                if (route_id, sequence) in seen:
                    raise ValueError("노선 안의 중복 순번")
                seen.add((route_id, sequence))
                route = routes.setdefault(route_id, {"routeId": route_id, "routeName": name, "stops": []})
                if route["routeName"] != name:
                    raise ValueError("같은 노선 ID의 이름 충돌")
                route["stops"].append({
                    "stopId": identifier(row[columns["NODE_ID"]]),
                    "stopName": identifier(row[columns["정류소명"]]),
                    "sequence": sequence, "longitude": longitude, "latitude": latitude,
                })
            except (ValueError, TypeError) as error:
                raise ValueError("Excel 행 {}: {}".format(row_number, error)) from error
        if not routes:
            raise ValueError("빈 버스 마스터")
        if route_ids and set(routes) != set(route_ids):
            raise ValueError("요청한 공식 노선 ID가 마스터에 없음")
        # 2. 원본 순번의 구멍은 누락 edge를 만들 수 있으므로 인접 구간을 생성하기 전에 차단한다.
        for route in routes.values():
            route["stops"].sort(key=lambda stop: stop["sequence"])
            if [stop["sequence"] for stop in route["stops"]] != list(range(1, len(route["stops"]) + 1)):
                raise ValueError("노선 {}의 순번 누락".format(route["routeId"]))
        return [routes[key] for key in sorted(routes)]
    finally:
        workbook.close()


def main():
    """출처 파일 해시를 기록하고 미검수 버스 전용 corpus를 새 파일에 저장한다.

    captured-at은 실제 수집 기록이 있을 때만 지정한다. 생략하면 null이며 평가기는 정적
    매핑만 수행한다. 시간 offset 검사는 Java 평가기에 맡긴다. 기존 출력은 덮어쓰지 않는다.
    """
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--response", type=Path, required=True)
    parser.add_argument("--bus-master", type=Path, required=True)
    parser.add_argument("--version", required=True)
    parser.add_argument("--case-id", required=True)
    parser.add_argument("--region", choices=["SEOUL", "OUTSIDE_SEOUL", "MIXED", "UNKNOWN"], required=True)
    parser.add_argument("--captured-at")
    parser.add_argument("--route-id", action="append", help="공식 노선 ID로 제한한 실험. 전체 정확도 측정에 사용 금지")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    raw = args.response.read_bytes()
    response = json.loads(raw.decode("utf-8-sig"))
    if response.get("status") != "OK" or not response.get("routes"):
        raise ValueError("비어 있지 않은 OK 경로 응답 필요")
    routes = read_bus_routes(args.bus_master, args.route_id)
    source = {
        "scope": "BUS_ONLY; 지하철 topology 미제공; labels 미검수",
        "busWorkbook": args.bus_master.name,
        "busWorkbookSha256": hashlib.sha256(args.bus_master.read_bytes()).hexdigest(),
        "responseFile": args.response.name,
        "responseSha256": hashlib.sha256(raw).hexdigest(),
        "selectedRouteIds": sorted(args.route_id) if args.route_id else "ALL",
    }
    corpus = {
        "schemaVersion": 1, "kind": "REAL", "masterDataVersion": args.version,
        "masterSource": json.dumps(source, ensure_ascii=False, sort_keys=True),
        "busMaxEndpointDistanceMeters": 150, "busRoutes": routes, "subwayLines": [],
        "aliases": {key: {} for key in ("busRouteAliasIds", "busStopAliasIds", "subwayLineAliasIds", "subwayStationAliasIds")},
        "captures": [{"caseId": args.case_id, "region": args.region,
                      "capturedAt": args.captured_at, "response": response, "labels": []}],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("x", encoding="utf-8") as stream:
        json.dump(corpus, stream, ensure_ascii=False, indent=2, allow_nan=False)
    print("버스 노선 {}개, 정류장 순번 {}개. 미검수 corpus 저장: {}".format(
        len(routes), sum(len(route["stops"]) for route in routes), args.output))


if __name__ == "__main__":
    main()
